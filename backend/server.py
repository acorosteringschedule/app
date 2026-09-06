from dotenv import load_dotenv
from pathlib import Path

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

import os
import io
import uuid
import logging
import calendar
from datetime import datetime, timezone, timedelta
from typing import List, Optional, Dict, Any

from fastapi import FastAPI, APIRouter, HTTPException, Depends, Request, UploadFile, File
from fastapi.responses import StreamingResponse
from fastapi.security import HTTPBearer, HTTPAuthorizationCredentials
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
from pydantic import BaseModel, Field, EmailStr
import bcrypt
import jwt
import openpyxl
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

# ---------- Setup ----------
mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

JWT_SECRET = os.environ['JWT_SECRET']
JWT_ALGO = "HS256"

app = FastAPI(title="ACO Shift Scheduler API")
api = APIRouter(prefix="/api")
security = HTTPBearer(auto_error=False)

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')
logger = logging.getLogger("aco")


def now_utc() -> datetime:
    return datetime.now(timezone.utc)


def new_id() -> str:
    return str(uuid.uuid4())


def hash_pw(pw: str) -> str:
    return bcrypt.hashpw(pw.encode(), bcrypt.gensalt()).decode()


def verify_pw(pw: str, hashed: str) -> bool:
    try:
        return bcrypt.checkpw(pw.encode(), hashed.encode())
    except Exception:
        return False


def make_token(user_id: str, role: str) -> str:
    payload = {"sub": user_id, "role": role, "exp": now_utc() + timedelta(days=7)}
    return jwt.encode(payload, JWT_SECRET, algorithm=JWT_ALGO)


async def get_current_user(request: Request, creds: Optional[HTTPAuthorizationCredentials] = Depends(security)) -> dict:
    token = None
    if creds:
        token = creds.credentials
    if not token:
        auth = request.headers.get("Authorization", "")
        if auth.startswith("Bearer "):
            token = auth[7:]
    if not token:
        raise HTTPException(401, "Not authenticated")
    try:
        payload = jwt.decode(token, JWT_SECRET, algorithms=[JWT_ALGO])
    except jwt.PyJWTError:
        raise HTTPException(401, "Invalid or expired token")
    user = await db.users.find_one({"id": payload["sub"]}, {"_id": 0, "password_hash": 0})
    if not user:
        raise HTTPException(401, "User not found")
    if user.get("status") != "approved" and user.get("role") != "admin":
        raise HTTPException(403, "Account not approved")
    return user


async def require_admin(user: dict = Depends(get_current_user)) -> dict:
    if user.get("role") != "admin":
        raise HTTPException(403, "Admin access required")
    return user


# ---------- Models ----------
class RegisterBody(BaseModel):
    nik: str
    name: str
    email: EmailStr
    password: str


class LoginBody(BaseModel):
    nik: str
    password: str


class PersonnelBody(BaseModel):
    nik: str
    name: str
    email: Optional[EmailStr] = None
    role: str = "personil"  # personil | admin
    order_index: int = 0
    avatar_base64: Optional[str] = None
    active: bool = True


class ShiftCellUpdate(BaseModel):
    user_id: str
    date: str  # YYYY-MM-DD
    shift: str  # pagi|siang|malam|off|cuti|sakit|dinas_luar|diklat|penugasan


class AutoGenerateBody(BaseModel):
    year: int
    month: int
    pattern: str  # "3-2" | "3-3" | "2-2"
    start_offsets: Optional[Dict[str, int]] = None  # user_id -> offset days


class RequestBody(BaseModel):
    type: str  # cuti_tahunan | cuti_penting | cuti_besar | sakit | dinas_luar | diklat | penugasan
    start_date: str
    end_date: str
    reason: Optional[str] = ""


class ApprovalBody(BaseModel):
    action: str  # approve | reject
    note: Optional[str] = ""


class SiteSettingsBody(BaseModel):
    title: Optional[str] = None
    subtitle: Optional[str] = None
    main_text: Optional[str] = None
    logo_base64: Optional[str] = None
    signature_base64: Optional[str] = None
    signature_name: Optional[str] = None
    primary_color: Optional[str] = None
    hero_image_base64: Optional[str] = None


# ---------- Startup ----------
@app.on_event("startup")
async def startup():
    await db.users.create_index("nik", unique=True)
    await db.users.create_index("email", unique=True)
    await db.shifts.create_index([("user_id", 1), ("date", 1)], unique=True)
    await db.requests.create_index("user_id")
    await db.change_logs.create_index("timestamp")

    # Seed admin
    admin_email = os.environ.get("ADMIN_EMAIL")
    admin_nik = os.environ.get("ADMIN_NIK", "ADMIN001")
    admin_pw = os.environ.get("ADMIN_PASSWORD", "Admin@123")
    admin_name = os.environ.get("ADMIN_NAME", "Administrator")
    existing = await db.users.find_one({"nik": admin_nik})
    if not existing:
        await db.users.insert_one({
            "id": new_id(),
            "nik": admin_nik,
            "name": admin_name,
            "email": admin_email,
            "password_hash": hash_pw(admin_pw),
            "role": "admin",
            "status": "approved",
            "order_index": -1,
            "active": True,
            "avatar_base64": None,
            "created_at": now_utc().isoformat(),
        })
        logger.info(f"Seeded admin: {admin_nik}")
    else:
        if not verify_pw(admin_pw, existing.get("password_hash", "")):
            await db.users.update_one({"nik": admin_nik}, {"$set": {"password_hash": hash_pw(admin_pw)}})

    # Seed default site settings
    settings = await db.site_settings.find_one({"id": "singleton"})
    if not settings:
        await db.site_settings.insert_one({
            "id": "singleton",
            "title": "ACO Shift Scheduler",
            "subtitle": "Sistem Penjadwalan Shift & Manajemen Personil",
            "main_text": "Kelola jadwal dinas, cuti, dan penugasan dengan presisi tinggi.",
            "logo_base64": None,
            "signature_base64": None,
            "signature_name": "",
            "hero_image_base64": None,
            "primary_color": "#008BFF",
            "updated_at": now_utc().isoformat(),
        })


# ---------- Auth Endpoints ----------
@api.post("/auth/register")
async def register(body: RegisterBody):
    exists = await db.users.find_one({"$or": [{"nik": body.nik}, {"email": body.email}]})
    if exists:
        raise HTTPException(400, "NIK atau email sudah terdaftar")
    max_order = await db.users.find_one({"role": "personil"}, sort=[("order_index", -1)])
    next_order = (max_order["order_index"] + 1) if max_order else 0
    user = {
        "id": new_id(),
        "nik": body.nik,
        "name": body.name,
        "email": body.email,
        "password_hash": hash_pw(body.password),
        "role": "personil",
        "status": "pending",
        "order_index": next_order,
        "active": True,
        "avatar_base64": None,
        "created_at": now_utc().isoformat(),
    }
    await db.users.insert_one(user)
    return {"message": "Pendaftaran berhasil. Menunggu persetujuan admin."}


@api.post("/auth/login")
async def login(body: LoginBody):
    user = await db.users.find_one({"nik": body.nik})
    if not user or not verify_pw(body.password, user["password_hash"]):
        raise HTTPException(401, "NIK atau password salah")
    if user["status"] == "pending":
        raise HTTPException(403, "Akun Anda belum disetujui admin")
    if user["status"] == "rejected":
        raise HTTPException(403, "Akun Anda ditolak")
    token = make_token(user["id"], user["role"])
    user.pop("password_hash", None)
    user.pop("_id", None)
    return {"token": token, "user": user}


@api.get("/auth/me")
async def me(user: dict = Depends(get_current_user)):
    return user


# ---------- User / Personnel Endpoints ----------
@api.get("/users")
async def list_users(user: dict = Depends(get_current_user), status: Optional[str] = None, include_admin: bool = False):
    q: Dict[str, Any] = {}
    if status:
        q["status"] = status
    if not include_admin:
        q["role"] = "personil"
    users = await db.users.find(q, {"_id": 0, "password_hash": 0}).sort("order_index", 1).to_list(500)
    return users


@api.get("/users/pending")
async def pending_users(admin: dict = Depends(require_admin)):
    users = await db.users.find({"status": "pending"}, {"_id": 0, "password_hash": 0}).to_list(200)
    return users


@api.post("/users/{user_id}/approve")
async def approve_user(user_id: str, body: ApprovalBody, admin: dict = Depends(require_admin)):
    target = await db.users.find_one({"id": user_id})
    if not target:
        raise HTTPException(404, "User tidak ditemukan")
    new_status = "approved" if body.action == "approve" else "rejected"
    await db.users.update_one({"id": user_id}, {"$set": {"status": new_status}})
    await db.notifications.insert_one({
        "id": new_id(),
        "user_id": user_id,
        "type": "account_" + new_status,
        "message": f"Akun Anda telah {'disetujui' if new_status == 'approved' else 'ditolak'} oleh admin.",
        "read": False,
        "created_at": now_utc().isoformat(),
    })
    await log_change(admin, f"{'Menyetujui' if new_status == 'approved' else 'Menolak'} akun {target['name']} ({target['nik']})")
    return {"status": new_status}


@api.post("/users")
async def create_user(body: PersonnelBody, admin: dict = Depends(require_admin)):
    exists = await db.users.find_one({"$or": [{"nik": body.nik}] + ([{"email": body.email}] if body.email else [])})
    if exists:
        raise HTTPException(400, "NIK atau email sudah terdaftar")
    max_order = await db.users.find_one({"role": "personil"}, sort=[("order_index", -1)])
    next_order = body.order_index if body.order_index else ((max_order["order_index"] + 1) if max_order else 0)
    default_pw = body.nik  # default password = NIK
    doc = {
        "id": new_id(),
        "nik": body.nik,
        "name": body.name,
        "email": body.email or f"{body.nik}@aco.local",
        "password_hash": hash_pw(default_pw),
        "role": body.role,
        "status": "approved",
        "order_index": next_order,
        "active": body.active,
        "avatar_base64": body.avatar_base64,
        "created_at": now_utc().isoformat(),
    }
    await db.users.insert_one(doc)
    await log_change(admin, f"Menambah personil {body.name} ({body.nik})")
    doc.pop("password_hash", None)
    doc.pop("_id", None)
    return doc


@api.put("/users/{user_id}")
async def update_user(user_id: str, body: PersonnelBody, admin: dict = Depends(require_admin)):
    upd = {k: v for k, v in body.model_dump().items() if v is not None}
    await db.users.update_one({"id": user_id}, {"$set": upd})
    await log_change(admin, f"Mengedit personil {body.name} ({body.nik})")
    updated = await db.users.find_one({"id": user_id}, {"_id": 0, "password_hash": 0})
    return updated


@api.delete("/users/{user_id}")
async def delete_user(user_id: str, admin: dict = Depends(require_admin)):
    target = await db.users.find_one({"id": user_id})
    if not target:
        raise HTTPException(404, "User tidak ditemukan")
    if target["role"] == "admin":
        raise HTTPException(400, "Tidak dapat menghapus admin")
    await db.users.delete_one({"id": user_id})
    await db.shifts.delete_many({"user_id": user_id})
    await log_change(admin, f"Menghapus personil {target['name']} ({target['nik']})")
    return {"ok": True}


@api.post("/users/reorder")
async def reorder_users(order: List[str], admin: dict = Depends(require_admin)):
    for idx, uid in enumerate(order):
        await db.users.update_one({"id": uid}, {"$set": {"order_index": idx}})
    await log_change(admin, "Mengubah urutan personil")
    return {"ok": True}


# ---------- Shift Schedule ----------
@api.get("/shifts")
async def get_shifts(year: int, month: int, user: dict = Depends(get_current_user)):
    _, ndays = calendar.monthrange(year, month)
    start = f"{year:04d}-{month:02d}-01"
    end = f"{year:04d}-{month:02d}-{ndays:02d}"
    shifts = await db.shifts.find({"date": {"$gte": start, "$lte": end}}, {"_id": 0}).to_list(5000)
    return {"year": year, "month": month, "days": ndays, "shifts": shifts}


@api.post("/shifts/cell")
async def upsert_cell(body: ShiftCellUpdate, admin: dict = Depends(require_admin)):
    existing = await db.shifts.find_one({"user_id": body.user_id, "date": body.date})
    prev = existing["shift"] if existing else None
    await db.shifts.update_one(
        {"user_id": body.user_id, "date": body.date},
        {"$set": {"user_id": body.user_id, "date": body.date, "shift": body.shift, "updated_at": now_utc().isoformat()}},
        upsert=True,
    )
    target = await db.users.find_one({"id": body.user_id}, {"_id": 0, "name": 1, "nik": 1})
    await log_change(admin, f"Ubah shift {target['name'] if target else body.user_id} pada {body.date}: {prev or '-'} → {body.shift}")
    return {"ok": True}


@api.post("/shifts/auto-generate")
async def auto_generate(body: AutoGenerateBody, admin: dict = Depends(require_admin)):
    _, ndays = calendar.monthrange(body.year, body.month)
    users = await db.users.find({"role": "personil", "active": True}).sort("order_index", 1).to_list(200)
    if not users:
        raise HTTPException(400, "Belum ada personil aktif")

    pattern_map = {
        "3-2": ["pagi", "siang", "malam", "off", "off"],
        "3-3": ["pagi", "siang", "malam", "off", "off", "off"],
        "2-2": ["pagi", "malam", "off", "off"],
    }
    seq = pattern_map.get(body.pattern)
    if not seq:
        raise HTTPException(400, "Pattern tidak valid")

    ops = []
    for i, u in enumerate(users):
        offset = (body.start_offsets or {}).get(u["id"], i % len(seq))
        for d in range(1, ndays + 1):
            shift = seq[(offset + d - 1) % len(seq)]
            date = f"{body.year:04d}-{body.month:02d}-{d:02d}"
            ops.append((u["id"], date, shift))

    for uid, date, shift in ops:
        await db.shifts.update_one(
            {"user_id": uid, "date": date},
            {"$set": {"user_id": uid, "date": date, "shift": shift, "updated_at": now_utc().isoformat()}},
            upsert=True,
        )
    await log_change(admin, f"Auto-generate jadwal {body.year}-{body.month:02d} pola {body.pattern}")
    return {"ok": True, "cells": len(ops)}


@api.post("/shifts/import-excel")
async def import_excel(year: int, month: int, file: UploadFile = File(...), admin: dict = Depends(require_admin)):
    content = await file.read()
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise HTTPException(400, "File kosong")
    header = rows[0]
    # Expected: [NIK, Nama, 1, 2, 3, ..., N]
    day_cols = []
    for idx, h in enumerate(header):
        if h is None:
            continue
        try:
            day = int(h)
            if 1 <= day <= 31:
                day_cols.append((idx, day))
        except (ValueError, TypeError):
            pass

    users_by_nik = {u["nik"]: u for u in await db.users.find({"role": "personil"}).to_list(500)}
    imported = 0
    valid_shifts = {"pagi", "siang", "malam", "off", "libur", "cuti", "sakit", "dinas_luar", "dl", "diklat", "penugasan"}
    for row in rows[1:]:
        if not row or not row[0]:
            continue
        nik = str(row[0]).strip()
        u = users_by_nik.get(nik)
        if not u:
            continue
        for col_idx, day in day_cols:
            if col_idx >= len(row):
                continue
            val = row[col_idx]
            if val is None:
                continue
            s = str(val).strip().lower().replace(" ", "_")
            if s == "libur":
                s = "off"
            if s == "dl":
                s = "dinas_luar"
            if s not in valid_shifts:
                continue
            date = f"{year:04d}-{month:02d}-{day:02d}"
            await db.shifts.update_one(
                {"user_id": u["id"], "date": date},
                {"$set": {"user_id": u["id"], "date": date, "shift": s, "updated_at": now_utc().isoformat()}},
                upsert=True,
            )
            imported += 1
    await log_change(admin, f"Import Excel jadwal {year}-{month:02d}: {imported} sel")
    return {"ok": True, "imported": imported}


# ---------- Requests (Cuti / Sakit / etc) ----------
def daterange(start: str, end: str):
    s = datetime.strptime(start, "%Y-%m-%d").date()
    e = datetime.strptime(end, "%Y-%m-%d").date()
    d = s
    while d <= e:
        yield d.isoformat()
        d += timedelta(days=1)


async def check_leave_conflict(exclude_user_id: str, start: str, end: str) -> List[str]:
    """Return list of dates that would violate minimum coverage (2 pagi + 2 siang + 1 malam)."""
    conflicts = []
    for date in daterange(start, end):
        pagi = await db.shifts.count_documents({"date": date, "shift": "pagi", "user_id": {"$ne": exclude_user_id}})
        siang = await db.shifts.count_documents({"date": date, "shift": "siang", "user_id": {"$ne": exclude_user_id}})
        malam = await db.shifts.count_documents({"date": date, "shift": "malam", "user_id": {"$ne": exclude_user_id}})
        current = await db.shifts.find_one({"date": date, "user_id": exclude_user_id})
        if current:
            cs = current["shift"]
            if cs == "pagi" and pagi < 2:
                conflicts.append(date)
                continue
            if cs == "siang" and siang < 2:
                conflicts.append(date)
                continue
            if cs == "malam" and malam < 1:
                conflicts.append(date)
                continue
    return conflicts


@api.post("/requests")
async def create_request(body: RequestBody, user: dict = Depends(get_current_user)):
    req = {
        "id": new_id(),
        "user_id": user["id"],
        "user_name": user["name"],
        "user_nik": user["nik"],
        "type": body.type,
        "start_date": body.start_date,
        "end_date": body.end_date,
        "reason": body.reason,
        "status": "pending",
        "note": "",
        "created_at": now_utc().isoformat(),
    }
    await db.requests.insert_one(req)
    # notify admins
    admins = await db.users.find({"role": "admin"}, {"_id": 0, "id": 1}).to_list(20)
    for a in admins:
        await db.notifications.insert_one({
            "id": new_id(),
            "user_id": a["id"],
            "type": "new_request",
            "message": f"Pengajuan baru dari {user['name']}: {body.type.replace('_', ' ').title()}",
            "ref_id": req["id"],
            "read": False,
            "created_at": now_utc().isoformat(),
        })
    req.pop("_id", None)
    return req


@api.get("/requests")
async def list_requests(user: dict = Depends(get_current_user), status: Optional[str] = None, mine: bool = False):
    q: Dict[str, Any] = {}
    if status:
        q["status"] = status
    if mine or user["role"] != "admin":
        q["user_id"] = user["id"]
    reqs = await db.requests.find(q, {"_id": 0}).sort("created_at", -1).to_list(500)
    return reqs


@api.post("/requests/{req_id}/action")
async def approve_request(req_id: str, body: ApprovalBody, admin: dict = Depends(require_admin)):
    req = await db.requests.find_one({"id": req_id})
    if not req:
        raise HTTPException(404, "Pengajuan tidak ditemukan")
    if body.action not in ("approve", "reject"):
        raise HTTPException(400, "Aksi tidak valid")

    if body.action == "approve" and req["type"] in ("cuti_tahunan", "cuti_penting", "cuti_besar", "sakit", "dinas_luar", "diklat", "penugasan"):
        conflicts = await check_leave_conflict(req["user_id"], req["start_date"], req["end_date"])
        if conflicts:
            await db.requests.update_one({"id": req_id}, {"$set": {
                "status": "rejected",
                "note": f"Ditolak otomatis: konflik minimum personil pada {', '.join(conflicts[:5])}",
                "decided_by": admin["name"],
                "decided_at": now_utc().isoformat(),
            }})
            await notify_request_result(req, "rejected", "Konflik minimum personil aktif")
            return {"status": "rejected", "reason": "conflict", "conflict_dates": conflicts}

    new_status = "approved" if body.action == "approve" else "rejected"
    await db.requests.update_one({"id": req_id}, {"$set": {
        "status": new_status,
        "note": body.note or "",
        "decided_by": admin["name"],
        "decided_at": now_utc().isoformat(),
    }})

    if new_status == "approved":
        # Update schedule cells accordingly
        type_to_shift = {
            "cuti_tahunan": "cuti", "cuti_penting": "cuti", "cuti_besar": "cuti",
            "sakit": "sakit", "dinas_luar": "dinas_luar", "diklat": "diklat", "penugasan": "penugasan",
        }
        cell_val = type_to_shift.get(req["type"], "cuti")
        for date in daterange(req["start_date"], req["end_date"]):
            await db.shifts.update_one(
                {"user_id": req["user_id"], "date": date},
                {"$set": {"user_id": req["user_id"], "date": date, "shift": cell_val, "updated_at": now_utc().isoformat()}},
                upsert=True,
            )

    await notify_request_result(req, new_status, body.note or "")
    await log_change(admin, f"{'Menyetujui' if new_status == 'approved' else 'Menolak'} pengajuan {req['type']} {req['user_name']} ({req['start_date']} - {req['end_date']})")
    return {"status": new_status}


async def notify_request_result(req: dict, status: str, note: str):
    msg = f"Pengajuan {req['type'].replace('_', ' ').title()} Anda telah {'disetujui' if status == 'approved' else 'ditolak'}"
    if note:
        msg += f". Catatan: {note}"
    await db.notifications.insert_one({
        "id": new_id(),
        "user_id": req["user_id"],
        "type": "request_" + status,
        "message": msg,
        "ref_id": req["id"],
        "read": False,
        "created_at": now_utc().isoformat(),
    })


# ---------- Notifications ----------
@api.get("/notifications")
async def get_notifications(user: dict = Depends(get_current_user)):
    notes = await db.notifications.find({"user_id": user["id"]}, {"_id": 0}).sort("created_at", -1).limit(50).to_list(50)
    return notes


@api.post("/notifications/read-all")
async def read_all(user: dict = Depends(get_current_user)):
    await db.notifications.update_many({"user_id": user["id"], "read": False}, {"$set": {"read": True}})
    return {"ok": True}


# ---------- Summary ----------
@api.get("/summary")
async def summary(year: int, month: int, user: dict = Depends(get_current_user)):
    _, ndays = calendar.monthrange(year, month)
    start = f"{year:04d}-{month:02d}-01"
    end = f"{year:04d}-{month:02d}-{ndays:02d}"
    users = await db.users.find({"role": "personil"}, {"_id": 0, "password_hash": 0}).sort("order_index", 1).to_list(500)
    shifts = await db.shifts.find({"date": {"$gte": start, "$lte": end}}, {"_id": 0}).to_list(5000)

    # hours per shift type
    hours_map = {"pagi": 6, "siang": 6, "malam": 12}
    result = []
    for u in users:
        counts = {"pagi": 0, "siang": 0, "malam": 0, "off": 0, "cuti": 0, "sakit": 0, "dinas_luar": 0, "diklat": 0, "penugasan": 0}
        for s in shifts:
            if s["user_id"] == u["id"]:
                if s["shift"] in counts:
                    counts[s["shift"]] += 1
        total_hours = counts["pagi"] * hours_map["pagi"] + counts["siang"] * hours_map["siang"] + counts["malam"] * hours_map["malam"]
        result.append({
            "user_id": u["id"],
            "nik": u["nik"],
            "name": u["name"],
            "counts": counts,
            "total_hours": total_hours,
        })
    return {"year": year, "month": month, "summary": result}


# ---------- Change Log ----------
async def log_change(actor: dict, message: str):
    await db.change_logs.insert_one({
        "id": new_id(),
        "actor_id": actor["id"],
        "actor_name": actor["name"],
        "message": message,
        "timestamp": now_utc().isoformat(),
    })


@api.get("/change-logs")
async def get_logs(admin: dict = Depends(require_admin), limit: int = 200):
    logs = await db.change_logs.find({}, {"_id": 0}).sort("timestamp", -1).limit(limit).to_list(limit)
    return logs


# ---------- Site Settings ----------
@api.get("/settings")
async def get_settings():
    s = await db.site_settings.find_one({"id": "singleton"}, {"_id": 0})
    return s or {}


@api.put("/settings")
async def update_settings(body: SiteSettingsBody, admin: dict = Depends(require_admin)):
    upd = {k: v for k, v in body.model_dump().items() if v is not None}
    upd["updated_at"] = now_utc().isoformat()
    await db.site_settings.update_one({"id": "singleton"}, {"$set": upd}, upsert=True)
    await log_change(admin, "Memperbarui pengaturan situs")
    s = await db.site_settings.find_one({"id": "singleton"}, {"_id": 0})
    return s


# ---------- Exports ----------
SHIFT_LABEL = {
    "pagi": "P", "siang": "S", "malam": "M", "off": "L",
    "cuti": "C", "sakit": "SK", "dinas_luar": "DL", "diklat": "DK", "penugasan": "PN",
}
SHIFT_HEX = {
    "pagi": colors.HexColor("#7DD3FC"), "siang": colors.HexColor("#FDE68A"), "malam": colors.HexColor("#C4B5FD"),
    "off": colors.HexColor("#E2E8F0"), "cuti": colors.HexColor("#FCA5A5"), "sakit": colors.HexColor("#FED7AA"),
    "dinas_luar": colors.HexColor("#86EFAC"), "diklat": colors.HexColor("#A5B4FC"), "penugasan": colors.HexColor("#F0ABFC"),
}


@api.get("/exports/pdf")
async def export_pdf(year: int, month: int, admin: dict = Depends(require_admin)):
    _, ndays = calendar.monthrange(year, month)
    start = f"{year:04d}-{month:02d}-01"
    end = f"{year:04d}-{month:02d}-{ndays:02d}"
    users = await db.users.find({"role": "personil", "active": True}, {"_id": 0}).sort("order_index", 1).to_list(500)
    shifts = await db.shifts.find({"date": {"$gte": start, "$lte": end}}, {"_id": 0}).to_list(5000)
    settings = await db.site_settings.find_one({"id": "singleton"}, {"_id": 0}) or {}

    shift_map: Dict[str, Dict[int, str]] = {}
    for s in shifts:
        uid = s["user_id"]
        day = int(s["date"].split("-")[2])
        shift_map.setdefault(uid, {})[day] = s["shift"]

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), leftMargin=8 * mm, rightMargin=8 * mm, topMargin=10 * mm, bottomMargin=10 * mm)
    styles = getSampleStyleSheet()
    story = []
    title_style = ParagraphStyle('t', parent=styles['Title'], fontSize=14, textColor=colors.HexColor("#0F172A"))
    sub_style = ParagraphStyle('s', parent=styles['Normal'], fontSize=9, textColor=colors.HexColor("#475569"))
    story.append(Paragraph(f"{settings.get('title', 'ACO Shift Scheduler')} — Jadwal Dinas", title_style))
    story.append(Paragraph(f"Periode: {calendar.month_name[month]} {year}", sub_style))
    story.append(Spacer(1, 4 * mm))

    header = ["No", "NIK", "Nama"] + [str(d) for d in range(1, ndays + 1)]
    data = [header]
    styles_list = [
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#0B0D14")),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 7),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('GRID', (0, 0), (-1, -1), 0.3, colors.HexColor("#94A3B8")),
    ]
    for i, u in enumerate(users):
        row = [str(i + 1), u["nik"], u["name"]]
        for d in range(1, ndays + 1):
            shift = shift_map.get(u["id"], {}).get(d, "")
            row.append(SHIFT_LABEL.get(shift, ""))
            if shift in SHIFT_HEX:
                styles_list.append(('BACKGROUND', (3 + d - 1, i + 1), (3 + d - 1, i + 1), SHIFT_HEX[shift]))
        data.append(row)

    col_widths = [10 * mm, 22 * mm, 40 * mm] + [7.2 * mm] * ndays
    t = Table(data, colWidths=col_widths, repeatRows=1)
    t.setStyle(TableStyle(styles_list))
    story.append(t)

    story.append(Spacer(1, 8 * mm))
    legend_style = ParagraphStyle('l', parent=styles['Normal'], fontSize=8)
    story.append(Paragraph("Keterangan: P=Pagi (07-13), S=Siang (13-19), M=Malam (19-07), L=Libur, C=Cuti, SK=Sakit, DL=Dinas Luar, DK=Diklat, PN=Penugasan", legend_style))

    if settings.get("signature_name"):
        story.append(Spacer(1, 15 * mm))
        story.append(Paragraph(f"Mengetahui,<br/><br/><br/><b>{settings.get('signature_name')}</b>", legend_style))

    doc.build(story)
    buffer.seek(0)
    return StreamingResponse(buffer, media_type="application/pdf", headers={"Content-Disposition": f'attachment; filename="jadwal_{year}_{month:02d}.pdf"'})


@api.get("/exports/xlsx")
async def export_xlsx(year: int, month: int, admin: dict = Depends(require_admin)):
    _, ndays = calendar.monthrange(year, month)
    start = f"{year:04d}-{month:02d}-01"
    end = f"{year:04d}-{month:02d}-{ndays:02d}"
    users = await db.users.find({"role": "personil", "active": True}, {"_id": 0}).sort("order_index", 1).to_list(500)
    shifts = await db.shifts.find({"date": {"$gte": start, "$lte": end}}, {"_id": 0}).to_list(5000)

    shift_map: Dict[str, Dict[int, str]] = {}
    for s in shifts:
        uid = s["user_id"]
        day = int(s["date"].split("-")[2])
        shift_map.setdefault(uid, {})[day] = s["shift"]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{year}-{month:02d}"
    ws.append(["NIK", "Nama"] + [i for i in range(1, ndays + 1)])
    for u in users:
        row = [u["nik"], u["name"]] + [SHIFT_LABEL.get(shift_map.get(u["id"], {}).get(d, ""), "") for d in range(1, ndays + 1)]
        ws.append(row)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return StreamingResponse(buffer, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="jadwal_{year}_{month:02d}.xlsx"'})


# ---------- Root ----------
@api.get("/")
async def root():
    return {"app": "ACO Shift Scheduler", "version": "1.0"}


app.include_router(api)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=os.environ.get('CORS_ORIGINS', '*').split(','),
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.on_event("shutdown")
async def shutdown():
    client.close()
